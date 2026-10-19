# -*- coding: utf-8 -*-
"""Ficha-cola (modelo aprovado em 12/10/2026) — 1 folha paisagem por leito.
Cabeçalho: leito + nome + Dx + dispositivos | alergia + plantão.
Corpo: CONTROLES + LAB | PRESCRIÇÃO (infusões/terapias + por sistemas) | PENDÊNCIAS + EXAMES + ANOTAÇÕES.
Rodapé de trabalho: grade de controles do plantão | intercorrências.
Preto no branco, Times New Roman, fios cinza entre linhas para guiar a leitura."""
import json, re, subprocess, sys, tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "skills" / "sasi-ingest-export" / "scripts"))
import build_passagem as motor  # noqa: E402
import estilo as E  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_TAB_ALIGNMENT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

FONTE = E.FONTE
S = {"tam": E.CORPO, "ls": E.ENTRELINHA_FICHA}
CINZA = E.PAPEL_CINZA
CFG = {"tam": E.CORPO}


def T(p, txt, b=False, tam=None, it=False):
    r = p.add_run(txt); r.bold = b; r.italic = it
    r.font.size = Pt(tam or CFG["tam"]); r.font.name = FONTE
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    return r


def rico(p, txt, tam=None, b_all=False):
    for i, parte in enumerate(re.split(r"\*\*", txt)):
        if parte:
            T(p, parte, b=b_all or i % 2 == 1, tam=tam)


def borda_p(p, lado="bottom", sz=6, cor=E.PAPEL_TINTA):
    pPr = p._p.get_or_add_pPr()
    b = pPr.find(qn("w:pBdr"))
    if b is None:
        b = OxmlElement("w:pBdr"); pPr.append(b)
    e = OxmlElement(f"w:{lado}")
    for k, v in (("val", "single"), ("sz", str(sz)), ("space", "1"), ("color", cor)):
        e.set(qn("w:" + k), v)
    b.append(e)


def tab(p, pos_cm, align="left"):
    p.paragraph_format.tab_stops.add_tab_stop(
        Cm(pos_cm), WD_TAB_ALIGNMENT.RIGHT if align == "right" else WD_TAB_ALIGNMENT.LEFT)


def tabela(doc_or_cell, rows, cols, larguras, bordas=None):
    t = doc_or_cell.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tb = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        v = (bordas or {}).get(lado)
        if v:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(v[0])); e.set(qn("w:color"), v[1])
        else:
            e.set(qn("w:val"), "nil")
        tb.append(e)
    tblPr.append(tb)
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(larguras[i] * 567)))
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = Cm(larguras[i])
    return t


def margens(cell, top=0, left=0, bottom=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)


def borda_cel(cell, **lados):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for lado, (sz, cor) in lados.items():
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), cor)
        b.append(e)
    tcPr.append(b)



W = E.LARGURA_UTIL
WL, WM = E.COL_CONTROLES, E.COL_PRESCRICAO
WR = W - WL - WM


def P(cell, first=False, after=1.5, before=0):
    if first and len(cell.paragraphs) == 1 and not cell.paragraphs[0].text:
        p = cell.paragraphs[0]
    else:
        p = cell.add_paragraph()
    f = p.paragraph_format
    f.space_before = Pt(before); f.space_after = Pt(after); f.line_spacing = S["ls"]
    return p


def t(p, txt, b=False, tam=None):
    return T(p, txt, b=b, tam=tam or S["tam"])


def rk(p, txt, tam=None):
    rico(p, txt, tam=tam or S["tam"])


def sec(cell, txt, first=False, before=5):
    p = P(cell, first=first, after=3, before=0 if first else before)
    p.paragraph_format.keep_with_next = True
    t(p, txt, b=True, tam=S["tam"])
    borda_p(p, "bottom", E.FIO_TITULO)
    return p


def subsec(cell, txt, before=5):
    p = P(cell, after=2, before=before)
    p.paragraph_format.keep_with_next = True
    t(p, txt, b=True, tam=S["tam"] - 1)
    return p


def encolher(p):
    """Parágrafo vazio que o Word exige depois de tabela dentro de célula: quase sem altura."""
    pPr = p._p.get_or_add_pPr()
    rPr = OxmlElement("w:rPr"); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "2"); rPr.append(sz)
    pPr.append(rPr)
    from docx.enum.text import WD_LINE_SPACING
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(1)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)


def grade(cell_or_doc, rows, cols, larg, linhas=True):
    bd = {"insideH": (E.FIO_GUIA, CINZA), "bottom": (E.FIO_GUIA, CINZA), "top": (E.FIO_GUIA, CINZA)} if linhas else None
    g = tabela(cell_or_doc, rows, cols, larg, bd)
    if hasattr(cell_or_doc, "paragraphs"):
        encolher(cell_or_doc.paragraphs[-1])
    for r in g.rows:
        for c in r.cells:
            margens(c, 10, 40, 10, 40)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return g


def droga(p, item):
    """Nome comercial/rotulo em negrito, resto normal. Itens com ** já marcados ficam como vieram."""
    if "**" in item:
        rk(p, item); return
    partes = item.split(" · ")
    for i, seg in enumerate(partes):
        if i:
            t(p, "  ·  ")
        m = re.match(r"^(.{2,28}?)(\s\(|\s\d|\s=>|:|$)(.*)$", seg)
        if m and not re.match(r"^\d", seg):
            t(p, m.group(1), b=True); t(p, m.group(2) + m.group(3))
        else:
            t(p, seg)


# ── dados do cabeçalho: uma fonte só ────────────────────────────────────────
# alergia, Dx e dispositivos moram no leito (alergia, diagnostico, dispositivos); o bloco `ficha` só sobrescreve
# quando precisa de um texto mais curto. Divergência entre os dois vira erro/aviso no gerador/validar.py.
CABECALHO = {"alergia": "alergia", "dx": "diagnostico", "disp": "dispositivos"}


def dado(leito, chave):
    v = leito["ficha"].get(chave)
    return v if v not in (None, "") else (leito.get(CABECALHO[chave]) or "")


def janelas(leito, meta):
    """Títulos de CONTROLES e BALANÇO. Sem `janela_controles`/`janela_balanco` no bloco, saem da janela do leito
    ou do plantão (24 h = 07 - 05 h)."""
    f = leito["ficha"]
    jan = leito.get("janela") or meta.get("janela") or "24 h"
    padrao_c = f"{meta.get('data', '')[:5]}  ·  07 - 05 h" if jan == "24 h" else jan
    padrao_b = "07 - 05 h" if jan == "24 h" else jan
    return f.get("janela_controles") or padrao_c, f.get("janela_balanco") or padrao_b


# ── colunas ────────────────────────────────────────────────────────────────
def linhas_controle(r, leito):
    jan_linhas = r["secoes"][0][1]
    out = []
    for l in jan_linhas:
        out += [x.strip() for x in l.split(" | ")] if " | " in l else [l]
    infus = set(leito.get("outros_balanco", []))
    vit, bal = {}, []
    for l in out:
        if l in infus:
            continue
        m = re.match(r"^([^:]{1,30}):\s*(.*)$", l)
        if not m:
            continue
        rot, resto = m.group(1), m.group(2)
        flags = re.findall(r"\[([^\]]+)\]", resto)
        val = re.sub(r"\s*\[[^\]]+\]", "", resto).strip()
        val = re.sub(r"\s*(mmHg|bpm|rpm|ºC|mg/dl)$", "", val)
        val = re.sub(r"\s*%(\s*\(.*\))?$", " %", val) if rot == "SpO2" else val
        flag = "  ".join(f.replace(" ", "") for f in flags)
        if rot in ("PAS", "PAD", "PAM", "FC", "FR", "SpO2", "TAX", "Dx"):
            vit[rot] = (val, flag)
        elif rot in ("Dieta", "Sup O2"):
            continue
        else:
            rot = {"Ingesta hídrica": "Ganhos", "Resíduo gástrico": "Resíduo", "UF (CVVHDF)": "UF",
                   "Débito SNG": "SNG"}.get(rot, rot)
            v = val.replace(" ml", "")
            if rot == "BH":
                v = v.replace("+ ", "+").replace("- ", "-")
            bal.append((rot, v))
    return vit, bal


def col_controles(C, r, leito, jan_c, jan_b):
    sec(C, f"CONTROLES  {jan_c}", first=True)
    vit, bal = linhas_controle(r, leito)
    pares = [("PAS", "FR"), ("PAD", "SpO2"), ("PAM", "TAX"), ("FC", None)]
    w = [1.0, 1.75, 1.3] * 2
    g = grade(C, len(pares), 6, w)
    for i, (a, b) in enumerate(pares):
        for j, k in enumerate((a, b)):
            if not k or k not in vit:
                continue
            val, fl = vit[k]
            c0, c1, c2 = g.rows[i].cells[j * 3: j * 3 + 3]
            t(P(c0, True, 0), k, b=True)
            t(P(c1, True, 0), val, b=bool(fl))
            if fl:
                q = P(c2, True, 0); q.alignment = 2; t(q, fl, b=True, tam=S["tam"] - 1)
    for row in g.rows:
        borda_cel(row.cells[3], left=(E.FIO_SECAO, CINZA))
    if "Dx" in vit:
        val, fl = vit["Dx"]
        g2 = grade(C, 1, 3, [1.05, 5.1, 2.15])
        t(P(g2.rows[0].cells[0], True, 0), "Dx", b=True)
        t(P(g2.rows[0].cells[1], True, 0), val, b=bool(fl))
        if fl:
            q = P(g2.rows[0].cells[2], True, 0); q.alignment = 2; t(q, fl, b=True, tam=S["tam"] - 1)

    subsec(C, f"BALANÇO  {jan_b}")
    curtos = [(k, v) for k, v in bal if k != "Evacuação"]
    evac = [v for k, v in bal if k == "Evacuação"]
    n = (len(curtos) + 1) // 2
    g3 = grade(C, n + (1 if evac else 0), 4, [1.55, 2.6, 1.55, 2.6])
    for i, (k, v) in enumerate(curtos):
        row, off = i // 2, (i % 2) * 2
        c0, c1 = g3.rows[row].cells[off], g3.rows[row].cells[off + 1]
        t(P(c0, True, 0), k, b=True)
        t(P(c1, True, 0), v + ("" if k in ("Evacuação",) else " ml"), b=(k == "BH"))
    for row in g3.rows[:n]:
        borda_cel(row.cells[2], left=(E.FIO_SECAO, CINZA))
    if evac:
        rr = g3.rows[n].cells
        t(P(rr[0], True, 0), "Evac.", b=True)
        m = rr[1].merge(rr[3]); t(P(m, True, 0), evac[0])

    sec(C, leito["ficha"].get("lab_titulo", "LABORATÓRIO"), before=8)
    for linha in leito["ficha"]["lab"]:
        q = P(C, after=2.5); rk(q, linha.replace(" | ", "   |   "))


def bloco_rx(C, linhas, w_rot=E.RX_ROTULO):
    g = grade(C, len(linhas), 2, [w_rot, WM - 0.6 - w_rot])
    for i, (rot, itens) in enumerate(linhas):
        c0, c1 = g.rows[i].cells
        c0.vertical_alignment = c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        t(P(c0, True, 0), rot, b=True, tam=S["tam"] - 1)
        for j, it in enumerate(itens):
            q = P(c1, first=(j == 0), after=0.5)
            if rot in ("VM", "O2", "TRS"):
                rk(q, it)
            else:
                droga(q, it)


def col_rx(C, leito, meta):
    f = leito["ficha"]
    sec(C, f"PRESCRIÇÃO  {meta['data_prescricao']}", first=True)
    subsec(C, "INFUSÕES E TERAPIAS", before=1)
    bloco_rx(C, f["infusoes"])
    subsec(C, "POR SISTEMAS", before=4)
    bloco_rx(C, f["sistemas"])


def col_acao(C, leito, n_anot):
    f = leito["ficha"]
    sec(C, "PENDÊNCIAS", first=True)
    for it in f["vigiar"]:
        q = P(C, after=2.5)
        q.paragraph_format.left_indent = Cm(0.55); q.paragraph_format.first_line_indent = Cm(-0.55)
        t(q, "☐  "); rk(q, it)
    if f.get("exames"):
        sec(C, "EXAMES")
        for it in f["exames"]:
            q = P(C, after=2)
            q.paragraph_format.left_indent = Cm(0.3); q.paragraph_format.first_line_indent = Cm(-0.3)
            t(q, "– "); rk(q, it)
    sec(C, "ANOTAÇÕES / NOVOS EXAMES")
    linhas(C, n_anot)


def linhas(cell, n):
    for k in range(n):
        q = P(cell, after=0, before=11)
        t(q, " ", tam=E.PAUTA); borda_p(q, "bottom", E.FIO_SECAO, E.PAPEL_CINZA_ANOTACAO)
        q.paragraph_format.right_indent = Cm(0.01 * (k % 2))


# ── página ─────────────────────────────────────────────────────────────────
def pagina(doc, meta, leito, n_anot, n_inter, primeiro):
    r = motor.compute_leito(leito, meta)
    if not primeiro:
        q = doc.add_paragraph(); q.paragraph_format.page_break_before = True
        q.paragraph_format.space_after = Pt(0); t(q, " ", tam=E.ESPACADOR)

    wbox, wdir = E.CAIXA_LEITO, E.CABECALHO_DIREITA
    h = tabela(doc, 1, 3, [wbox, W - wbox - wdir, wdir])
    a, b, c = h.rows[0].cells
    for cc in (a, b, c):
        cc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    caixa = (E.FIO_CAIXA, E.PAPEL_TINTA)
    borda_cel(a, top=caixa, left=caixa, bottom=caixa, right=caixa)
    margens(a, 20, 40, 20, 40); margens(b, 0, 200, 0, 100); margens(c, 0, 100, 0, 0)
    q = P(a, True, 0); q.alignment = 1; t(q, "LEITO", b=True, tam=E.LEITO_ROTULO)
    q = P(a, after=0); q.alignment = 1; t(q, leito["leito"], b=True, tam=E.LEITO)
    nome = leito["nome"].title().replace(" Da ", " da ").replace(" De ", " de ")
    q = P(b, True, 1)
    t(q, nome, b=True, tam=E.NOME)
    t(q, f"     {leito['idade_sexo']}  ·  {leito['dia_uti'].split(' (')[0]}", tam=S["tam"])
    q = P(b, after=1); q.paragraph_format.left_indent = Cm(0.75); q.paragraph_format.first_line_indent = Cm(-0.75)
    t(q, "Dx:  ", b=True); rk(q, dado(leito, "dx"))
    q = P(b, after=0); q.paragraph_format.left_indent = Cm(0.75); q.paragraph_format.first_line_indent = Cm(-0.75)
    t(q, "Disp.:  ", b=True); rk(q, dado(leito, "disp"))
    q = P(c, True, 2); t(q, "Alergia:  ", b=True)
    al = dado(leito, "alergia")
    t(q, al.upper() if al.lower() != "nega" else "nega", b=al.lower() != "nega")
    q = P(c, after=0); t(q, f"{meta['turno'].capitalize()}  ·  {meta['data']}")
    q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(4); q.paragraph_format.space_before = Pt(2)
    t(q, " ", tam=E.ESPACADOR); borda_p(q, "bottom", E.FIO_CAIXA)

    body = tabela(doc, 1, 3, [WL, WM, WR], {"insideV": (E.FIO_SECAO, E.PAPEL_TINTA)})
    esq, meio, dir_ = body.rows[0].cells
    margens(esq, 0, 0, 0, 150); margens(meio, 0, 150, 0, 150); margens(dir_, 0, 150, 0, 0)
    col_controles(esq, r, leito, *janelas(leito, meta))
    col_rx(meio, leito, meta)
    col_acao(dir_, leito, n_anot)

    q = doc.add_paragraph(); q.paragraph_format.space_before = Pt(6); q.paragraph_format.space_after = Pt(4)
    t(q, " ", tam=E.ESPACADOR); borda_p(q, "bottom", E.FIO_TITULO)
    wl = E.FAIXA_CONTROLES
    bt = tabela(doc, 1, 2, [wl, W - wl], {"insideV": (E.FIO_SECAO, E.PAPEL_TINTA)})
    L, R = bt.rows[0].cells
    margens(L, 0, 0, 0, 150); margens(R, 0, 150, 0, 0)
    sec(L, "CONTROLES DO PLANTÃO", first=True)
    g = tabela(L, 6, 9, [E.GRADE_ROTULO] + [E.GRADE_CELULA] * 8,
               {k: (E.FIO_SECAO, E.PAPEL_TINTA) for k in ("top", "left", "bottom", "right", "insideH", "insideV")})
    encolher(L.paragraphs[-1])
    for i, rot in enumerate(("Hora", "PA (PAM)", "FC", "SpO2 / FiO2", "Diurese", "Dx")):
        cel = g.rows[i].cells[0]; margens(cel, 15, 60, 15, 40)
        t(P(cel, True, 0), rot, b=True, tam=S["tam"] - 1)
        for j in range(1, 9):
            t(P(g.rows[i].cells[j], True, 0), " ", tam=S["tam"] + 1)
    sec(R, "INTERCORRÊNCIAS, CONDUTAS E PASSAGEM", first=True)
    linhas(R, n_inter)


ORDEM = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
           "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection "
           "textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark "
            "headers cellIns cellDel cellMerge tcPrChange",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
             "tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription tblPrChange",
}


def sanear(doc):
    """Reordena pPr/tcPr/tblPr na ordem do schema OOXML (o Word é estrito; o LibreOffice não)."""
    for tag, ordem in ORDEM.items():
        pos = {qn("w:" + n): i for i, n in enumerate(ordem.split())}
        for el in doc.element.body.iter(qn("w:" + tag)):
            filhos = list(el)
            vistos, unicos = set(), []
            for f in filhos:
                if f.tag in vistos and f.tag != qn("w:tab"):
                    el.remove(f); continue
                vistos.add(f.tag); unicos.append(f)
            for f in unicos:
                el.remove(f)
            for f in sorted(unicos, key=lambda x: pos.get(x.tag, 999)):
                el.append(f)


def gerar(plantao, caminho, cfg=None, folga=0.0):
    meta = plantao["meta"]
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = FONTE; st.font.size = Pt(E.CORPO)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.space_before = Pt(0)
    s = doc.sections[0]; s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(E.A4_LONGO), Cm(E.A4_CURTO)
    s.left_margin = s.right_margin = Cm(E.MARGEM_FICHA_LATERAL); s.top_margin = Cm(E.MARGEM_FICHA_TOPO)
    s.bottom_margin = Cm(E.MARGEM_FICHA_PE + folga)
    s.header_distance = s.footer_distance = Cm(E.DISTANCIA_CABECALHO)
    for i, leito in enumerate(plantao["leitos"]):
        n_anot, n_inter, tam = (cfg or {}).get(leito["leito"], (6, 4, E.CORPO))
        S["tam"] = tam
        pagina(doc, meta, leito, n_anot, n_inter, i == 0)
    body = doc.element.body
    if body[0].tag == qn("w:p") and not "".join(body[0].itertext()).strip():
        body.remove(body[0])
    S["tam"] = E.CORPO
    sanear(doc)
    doc.save(caminho)


def n_paginas(docx):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(Path(docx).parent), str(docx)],
                   check=True, capture_output=True, timeout=180)
    out = subprocess.run(["pdfinfo", str(Path(docx).with_suffix(".pdf"))], capture_output=True, text=True).stdout
    return int(re.search(r"Pages:\s+(\d+)", out).group(1))


def ajustar(plantao, folga=E.FOLGA_MEDICAO):
    """Devolve (cfg por leito, leitos que não couberam nem em 9 pt sem anotações)."""
    res, nao_coube = {}, []
    with tempfile.TemporaryDirectory() as tmp:
        for leito in plantao["leitos"]:
            um = {"meta": plantao["meta"], "leitos": [leito]}
            achou = None
            for tam in E.AJUSTE:
                for n_anot, n_inter in ((14, 5), (12, 5), (10, 5), (9, 4), (8, 4), (7, 4), (6, 4), (5, 4),
                                        (4, 3), (3, 3), (2, 3), (1, 3), (0, 3)):
                    fx = Path(tmp) / "x.docx"
                    gerar(um, fx, {leito["leito"]: (n_anot, n_inter, tam)}, folga=folga)
                    if n_paginas(fx) == 1:
                        achou = (n_anot, n_inter, tam); break
                if achou:
                    break
            res[leito["leito"]] = achou or (0, 2, 9)
            if not achou:
                nao_coube.append(leito["leito"])
            print(leito["leito"], res[leito["leito"]], "OK" if achou else "NAO COUBE", file=sys.stderr)
    return res, nao_coube


if __name__ == "__main__":
    pl = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    cfg, nao_coube = ajustar(pl)
    gerar(pl, sys.argv[2], cfg)
    print(sys.argv[2])
    if nao_coube:
        sys.exit("NAO COUBE em 1 folha: leito(s) " + ", ".join(nao_coube))
