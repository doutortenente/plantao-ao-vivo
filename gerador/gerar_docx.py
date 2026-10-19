#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gerar_docx.py — gera os dois Word do plantão a partir de UM arquivo JSON.

  1. DD-MM-AA-utiN-evolucoes-<turno>.docx   (retrato, 1 leito por página)
     Só os blocos que mudam no dia, na ordem do docs/BRIEFING.md.
  2. DD-MM-AA-utiN-ficha-cola-<turno>.docx  (paisagem, 1 FOLHA por leito) + .pdf
     Modelo aprovado em 12/10/2026: gerador/ficha_cola.py, a partir do bloco `ficha` de cada leito.

Fonte Times New Roman; destaque só por negrito (rótulos, flags, alertas). Sem cor.
Toda conta (Máx - Mín, flags, balanço) vem do motor:
skills/sasi-ingest-export/scripts/build_passagem.py

Uso:
  python3 gerador/gerar_docx.py plantao.json --saida pasta/
  python3 gerador/gerar_docx.py plantao.json --saida pasta/ --sem-ajuste   # não mede no LibreOffice

O ajuste (padrão) mede cada leito no LibreOffice para a ficha caber em 1 folha. Sem ajuste, a ficha sai com
6 linhas de anotação e 4 de intercorrências em 10 pt, e pode estourar.
Antes de gerar roda gerador/validar.py (erro = não gera). No fim roda gerador/conferir.py (fonte, tamanho, cor, texto proibido, 1 folha por leito, PNG de cada página).
Saída: 0 = ok; 2 = JSON sem bloco `ficha` ou sem LibreOffice; 3 = algum leito não coube em 1 folha;
4 = a conferência achou erro.
Contrato do JSON: gerador/schema-plantao.md
"""
import argparse
import json
import re
import shutil
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(RAIZ / "skills" / "sasi-ingest-export" / "scripts"))
import build_passagem as motor  # noqa: E402
import estilo as E  # noqa: E402

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

FONTE = E.FONTE
TAM = E.CORPO
FLAG = re.compile(r"(\[[^\]]*x [<>][^\]]*\])")


# ── utilidades de formatação ─────────────────────────────────────────────────
def base(doc, entrelinha=E.ENTRELINHA_EVOLUCAO):
    st = doc.styles["Normal"]
    st.font.name = FONTE
    st.font.size = Pt(TAM)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    pf = st.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = entrelinha


def T(p, texto, negrito=False, tam=None):
    r = p.add_run(texto)
    r.bold = negrito
    r.font.size = Pt(tam or TAM)
    return r


def com_flags(p, texto, tam=None):
    for parte in FLAG.split(texto):
        if parte:
            if FLAG.fullmatch(parte):
                T(p, parte, True, tam)
            else:
                T(p, parte, False, tam)


def filete(p, sz=6, cor=E.PAPEL_TINTA):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", str(sz)), ("space", "1"), ("color", cor)):
        e.set(qn("w:" + k), v)
    b.append(e); pPr.append(b)


def rotulo_valor(p, linha, tam=None):
    m = re.match(r"^(-\s)?([^:]{1,40}):(.*)$", linha)
    if m:
        T(p, (m.group(1) or "") + m.group(2) + ":", True, tam)
        com_flags(p, m.group(3), tam)
    else:
        com_flags(p, linha, tam)


# ── EVOLUÇÕES ────────────────────────────────────────────────────────────────
def gerar_evolucoes(plantao, caminho):
    meta = plantao.get("meta", {})
    doc = Document(); base(doc)
    s = doc.sections[0]
    s.page_height, s.page_width = Cm(E.A4_LONGO), Cm(E.A4_CURTO)
    s.left_margin = s.right_margin = Cm(E.MARGEM_EVOLUCAO_LATERAL)
    s.top_margin = s.bottom_margin = Cm(E.MARGEM_EVOLUCAO_VERTICAL)
    for i, leito in enumerate(plantao["leitos"]):
        r = motor.compute_leito(leito, meta)
        p = doc.add_paragraph()
        if i:
            p.paragraph_format.page_break_before = True
        T(p, f"LEITO {leito['leito']}  ", True, tam=TAM)
        T(p, leito.get("nome", ""), True, tam=TAM)
        p2 = doc.add_paragraph()
        info = [leito.get("idade_sexo"), leito.get("dia_uti"),
                f"Evolução {({'noturno': 'noturna', 'diurno': 'diurna'}).get(meta.get('turno', ''), meta.get('turno', ''))} {meta.get('unidade', '')}".strip(), meta.get("data")]
        T(p2, "  ·  ".join(x for x in info if x))
        filete(p2, E.FIO_TITULO); p2.paragraph_format.space_after = Pt(6)
        for titulo, linhas in r["secoes"]:
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(9); h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            T(h, titulo, True, tam=TAM); filete(h, E.FIO_SECAO, E.PAPEL_CINZA_SECAO)
            plano = titulo.startswith("Plano")
            for l in linhas:
                if not l.strip():
                    continue
                q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(1.5)
                if re.match(r"^(\d+\.\s|-\s)", l):
                    q.paragraph_format.left_indent = Cm(0.5); q.paragraph_format.first_line_indent = Cm(-0.5)
                if plano and l.startswith("- "):
                    q.paragraph_format.left_indent = Cm(0.9); q.paragraph_format.first_line_indent = Cm(-0.4)
                m_num = re.match(r"^(\d+\.\s[^:]{1,40}:)(.*)$", l) if plano else re.match(r"^(\d+\.\s)(.*)$", l)
                if titulo.startswith(("Evolução",)):
                    com_flags(q, l)
                elif m_num and (plano or titulo.startswith("Impressão")):
                    T(q, m_num.group(1), True); T(q, m_num.group(2))
                elif plano or titulo.startswith("Impressão"):
                    T(q, l)
                else:
                    rotulo_valor(q, l)
    doc.save(caminho)


def nome_arquivo(meta, classe, tipo):
    """Padrão DD-MM-AA-utiN-<classe>.<tipo>, ex.: 11-10-26-uti2-ficha-cola-diurno.docx"""
    d = re.match(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", meta.get("data", ""))
    data = f"{int(d.group(1)):02d}-{int(d.group(2)):02d}-{d.group(3)[-2:]}" if d else "sem-data"
    uti = re.sub(r"\s+", "", meta.get("unidade", "uti")).lower()
    return f"{data}-{uti}-{classe}.{tipo}"


def main():
    ap = argparse.ArgumentParser(description="Gera evolução + ficha-cola (.docx) do plantão.")
    ap.add_argument("json")
    ap.add_argument("--saida", default="saida")
    ap.add_argument("--sem-ajuste", action="store_true")
    a = ap.parse_args()
    plantao = json.loads(Path(a.json).read_text(encoding="utf-8"))
    meta = plantao.get("meta", {})

    sem_ficha = [l.get("leito", "?") for l in plantao["leitos"] if "ficha" not in l]
    if sem_ficha:
        print("ERRO: leito(s) " + ", ".join(sem_ficha) + " sem o bloco `ficha` (gerador/schema-plantao.md). "
              "A ficha-cola antiga foi aposentada; monte o bloco e gere de novo.", file=sys.stderr)
        sys.exit(2)
    plantao_original = plantao
    plantao, _ = motor.preparar(plantao)  # dias de ATB (início = D0) e de UTI (admissão = D1)
    import validar
    erros_v, avisos_v = validar.validar(plantao)
    if erros_v:
        print("ERRO: o JSON não passou na validação (gerador/validar.py):", file=sys.stderr)
        for e in erros_v:
            print("- " + e, file=sys.stderr)
        sys.exit(2)
    soffice = shutil.which("soffice")
    if not a.sem_ajuste and not (soffice and shutil.which("pdfinfo")):
        print("ERRO: o ajuste da ficha-cola precisa do LibreOffice (soffice) e do pdfinfo (poppler-utils). "
              "Instale os dois ou rode com --sem-ajuste (a folha pode estourar).", file=sys.stderr)
        sys.exit(2)

    pasta = Path(a.saida); pasta.mkdir(parents=True, exist_ok=True)
    turno = re.sub(r"\s+", "-", meta.get("turno", "").strip().lower())
    sufixo = f"-{turno}" if turno else ""
    evo = pasta / nome_arquivo(meta, "evolucoes" + sufixo, "docx")
    cola = pasta / nome_arquivo(meta, "ficha-cola" + sufixo, "docx")
    gerar_evolucoes(plantao, evo)

    import ficha_cola
    cfg, nao_coube = ({}, []) if a.sem_ajuste else ficha_cola.ajustar(plantao)
    ficha_cola.gerar(plantao, cola, cfg)
    if soffice:
        import office
        office.para_pdf(cola, pasta)
    else:
        print("AVISO: sem LibreOffice, o PDF da ficha-cola não foi gerado.", file=sys.stderr)
    print(evo); print(cola)
    _, avisos = motor.build(plantao_original)
    avisos = avisos_v + avisos
    if avisos:
        print("\nRodapé (conferir — fora do bloco):")
        for w in avisos:
            print("- " + w)
    import conferir
    erros, avisos_conf, pngs = conferir.conferir(evo, cola, len(plantao["leitos"]))
    print(conferir.relatorio(erros, avisos_conf, pngs))
    if nao_coube:
        print("\nERRO: leito(s) " + ", ".join(nao_coube) + " não couberam em 1 folha nem em 9 pt sem anotações. "
              "Encurte a prescrição desses leitos e gere de novo.", file=sys.stderr)
        sys.exit(3)
    if erros:
        sys.exit(4)


if __name__ == "__main__":
    main()
