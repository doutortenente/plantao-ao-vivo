#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gerar_admissao.py — Word da nota de admissão (skills/admissao-uti), no padrão das evoluções reais do Dr.

A4 (decisão de 19/10/2026; antes era Carta), margens esq. 1,25, dir. 1,34, sup. e inf. 2,5 cm, Times New Roman 10,
espaçamento simples. Negrito só nas duas linhas de cabeçalho, na frase "Paciente de...", nos títulos de seção, nos
rótulos em linha, nos rótulos dos sinais vitais, no nome do exame complementar e nos sistemas numerados do plano.
O plano não se parte entre páginas.

Uso:
  python3 gerador/gerar_admissao.py nota.txt saida/            # nome DD-MM-AA-utiN-admissao-leitoLL.docx (do cabeçalho)
  python3 gerador/gerar_admissao.py nota.txt saida/x.docx      # nome escolhido
Depois de gerar: PDF, nº de páginas e a conferência (fonte, tamanho, cor, "(revisar)", "?" e reticências).
"""
import argparse
import re
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import estilo as E  # noqa: E402

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

PLANO = "Plano terapêutico e Condutas:"
SECOES = ("Dispositivos:", "Uso:", "Sinais Vitais admissão:", "Exame Físico Geral:",
          "Exames Complementares:", PLANO)
EM_LINHA = re.compile(r"^(Admissão \([^)]*\):|HPMA:|Antecedentes:|Medicamentos de uso domiciliar:|ALERGIAS:)")


def gerar(txt, saida, margem=E.MARGEM_ADMISSAO_VERTICAL):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = E.FONTE; st.font.size = Pt(E.CORPO)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), E.FONTE)
    pf = st.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = E.ENTRELINHA_ADMISSAO
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(E.A4_CURTO), Cm(E.A4_LONGO)
    s.left_margin, s.right_margin = Cm(E.MARGEM_ADMISSAO_ESQ), Cm(E.MARGEM_ADMISSAO_DIR)
    s.top_margin = s.bottom_margin = Cm(margem)
    secao, cab = None, 0
    for linha in Path(txt).read_text(encoding="utf-8").rstrip().split("\n"):
        p = doc.add_paragraph()
        if secao == PLANO:
            p.paragraph_format.keep_with_next = True  # plano não se parte entre páginas
        if not linha.strip():
            continue

        def run(t, negrito=False):
            r = p.add_run(t); r.bold = negrito
            r.font.name = E.FONTE; r.font.size = Pt(E.CORPO)

        if linha.strip() in SECOES:
            secao = linha.strip(); run(linha.strip(), True)
            if secao == PLANO:
                p.paragraph_format.keep_with_next = True
            continue
        if cab < 2:  # "LL - NOME" e "ADMISSÃO UTI N - DD-MM-AA"
            cab += 1; run(linha, True); continue
        if linha.startswith("Paciente de"):
            run(linha, True); continue
        m = None
        if secao is None:
            m = EM_LINHA.match(linha)
        elif secao == "Sinais Vitais admissão:":
            m = re.match(r"^[^:]{1,10}:", linha)
        elif secao == "Exames Complementares:" and linha.startswith("- "):
            m = re.match(r"^- [^:]{1,80}:", linha)
        elif secao == PLANO:
            m = re.match(r"^\d+\. [^:]{1,40}:", linha)
            if not m and re.match(r"^\d+\. ", linha):
                run(linha, True); continue
        if m:
            run(m.group(0), True); run(linha[m.end():])
        else:
            run(linha)
    doc.save(saida)


def nome_do_cabecalho(txt):
    """'02 - NOME' + 'ADMISSÃO UTI 2 - 10-10-26' -> '10-10-26-uti2-admissao-leito02.docx'."""
    linhas = [l for l in Path(txt).read_text(encoding="utf-8").splitlines() if l.strip()][:2]
    m1 = re.match(r"^\s*(\d{1,3})\s*[-–—]", linhas[0]) if linhas else None
    m2 = re.search(r"UTI\s*(\d+)\s*[-–—]\s*(\d{2})-(\d{2})-(\d{2,4})", linhas[1]) if len(linhas) > 1 else None
    if not (m1 and m2):
        return None
    return f"{m2.group(2)}-{m2.group(3)}-{m2.group(4)[-2:]}-uti{m2.group(1)}-admissao-leito{m1.group(1).zfill(2)}.docx"


def conferir_admissao(docx):
    """Erros de fonte/tamanho/cor/texto + nº de páginas (ou None sem LibreOffice)."""
    import conferir
    erros, _ = conferir.conferir_docx(docx, {E.CORPO})
    for txt in conferir._texto_paragrafos(Document(str(docx))):
        if re.search(r"(?<![\w])\?(?![\w])|\.\.\.|…|\[ \]|não informado", txt):
            erros.append(f"{Path(docx).name}: na admissão, campo sem dado fica só com o rótulo: {txt[:80]}")
    n = None
    if (shutil.which("soffice") or shutil.which("libreoffice")) and shutil.which("pdfinfo"):
        import office
        n = conferir.paginas(office.para_pdf(docx))
    return erros, n


def main():
    ap = argparse.ArgumentParser(description="Gera o Word da nota de admissão.")
    ap.add_argument("nota"); ap.add_argument("saida")
    ap.add_argument("--margem", type=float, default=E.MARGEM_ADMISSAO_VERTICAL, help="sup./inf. em cm")
    a = ap.parse_args()
    destino = Path(a.saida)
    if destino.suffix.lower() != ".docx":
        nome = nome_do_cabecalho(a.nota)
        if not nome:
            sys.exit("ERRO: cabeçalho fora do modelo ('LL - NOME' / 'ADMISSÃO UTI N - DD-MM-AA'); passe o nome .docx")
        destino.mkdir(parents=True, exist_ok=True)
        destino = destino / nome
    gerar(a.nota, destino, a.margem)
    erros, n = conferir_admissao(destino)
    print(destino)
    if n:
        print(f"{n} página(s)" + (" — se passou de 1, o plano tem de estar inteiro na 2ª" if n > 1 else ""))
    for e in erros:
        print("ERRO: " + e)
    sys.exit(4 if erros else 0)


if __name__ == "__main__":
    main()
