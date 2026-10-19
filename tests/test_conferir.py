# -*- coding: utf-8 -*-
"""Testes do gerador/conferir.py. Rodar: python3 -m unittest discover -s tests"""
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "gerador"))
import conferir as C  # noqa: E402
import estilo as E  # noqa: E402

from docx import Document  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

TEM_PDF = bool(shutil.which("soffice") and shutil.which("pdfinfo"))


def docx_com(tmp, linhas, fonte=E.FONTE, tam=E.CORPO, cor=None):
    doc = Document()
    doc.styles["Normal"].font.name = E.FONTE
    for txt in linhas:
        r = doc.add_paragraph().add_run(txt)
        r.font.name = fonte; r.font.size = Pt(tam)
        if cor:
            r.font.color.rgb = RGBColor(0xC0, 0, 0)
    caminho = Path(tmp) / "x.docx"
    doc.save(caminho)
    return caminho


class Docx(unittest.TestCase):
    def test_limpo_passa(self):
        with tempfile.TemporaryDirectory() as tmp:
            erros, avisos = C.conferir_docx(docx_com(tmp, ["PAS: 132 - 88 mmHg [1x < 90]", "– RX tórax"]),
                                            C.TAM_FICHA)
        self.assertEqual(erros, [])

    def test_pega_revisar_travessao_seta(self):
        with tempfile.TemporaryDirectory() as tmp:
            erros, _ = C.conferir_docx(docx_com(tmp, ["PAS: 132 (revisar) - 88", "FC: 112 – 88 bpm",
                                                      "1. IRA ↑ CR 1,7"]), C.TAM_EVOLUCAO, evolucao=True)
        txt = "\n".join(erros)
        self.assertIn("(revisar)", txt)
        self.assertIn("travessão", txt)
        self.assertIn("seta", txt)

    def test_pega_cor_fonte_tamanho(self):
        with tempfile.TemporaryDirectory() as tmp:
            erros, _ = C.conferir_docx(docx_com(tmp, ["K 3,3"], fonte="Arial", tam=12, cor=True), C.TAM_EVOLUCAO)
        txt = "\n".join(erros)
        self.assertIn("Arial", txt)
        self.assertIn("12 pt", txt)
        self.assertIn("cor", txt)

    def test_excecoes_aprovadas_na_ficha(self):
        for t in (8, 9, 13, 20, 9.5, 8.5):
            self.assertIn(t, C.TAM_FICHA)
        self.assertNotIn(13, C.TAM_EVOLUCAO)


@unittest.skipUnless(TEM_PDF, "sem LibreOffice/pdfinfo")
class Paginas(unittest.TestCase):
    def test_conta_folhas_por_leito(self):
        sys.path.insert(0, str(RAIZ / "gerador"))
        import gerar_docx
        import ficha_cola
        dados = json.loads((RAIZ / "gerador" / "exemplo" / "plantao-exemplo.json").read_text(encoding="utf-8"))
        with tempfile.TemporaryDirectory() as tmp:
            evo, cola = Path(tmp) / "evo.docx", Path(tmp) / "cola.docx"
            gerar_docx.gerar_evolucoes(dados, evo)
            ficha_cola.gerar(dados, cola, {})
            erros, _, pngs = C.conferir(evo, cola, 1)
            self.assertEqual(erros, [])
            self.assertTrue(pngs)
            erros, _, _ = C.conferir(evo, cola, 2)
            self.assertTrue(any("1 folha por leito" in e for e in erros))


if __name__ == "__main__":
    unittest.main()
